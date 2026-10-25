"""Entirely fictional Fuel Store-style workbooks, generated in memory."""

from datetime import datetime
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook
from openpyxl.utils.datetime import CALENDAR_MAC_1904


HEADERS = (
    "Cost Centre", "Card No", "Date", "Card VRN", "Driver", "Tran VRN", "Odo",
    "Card Type", "Product", "Site No", "Site Name", "Quantity", "Unit Price", "Cost",
)

TRANSACTIONS = (
    {
        "Cost Centre": "Fictional depot", "Card No": "FICTIONAL-CARD-0001",
        "Date": datetime(2027, 2, 9, 6, 12, 34), "Card VRN": "ZZ26 ABC",
        "Driver": "Fictional Driver One", "Tran VRN": "ZZ26 ABC", "Odo": 123456,
        "Card Type": "FUEL STORE", "Product": "KEYFUELS DIESEL", "Site No": "FAKE-01",
        "Site Name": "Fictional service station", "Quantity": 143.27,
        "Unit Price": 1.30, "Cost": 186.25,
    },
    {
        "Cost Centre": "Fictional depot", "Card No": "FICTIONAL-CARD-0002",
        "Date": datetime(2027, 2, 10, 17, 45), "Card VRN": "ZX26 DEF",
        "Driver": "Fictional Driver Two", "Tran VRN": "ZX26 DEF", "Odo": 234567,
        "Card Type": "UK FUELS", "Product": "UK FUELS DIESEL", "Site No": "FAKE-02",
        "Site Name": "Another fictional station", "Quantity": 211.08,
        "Unit Price": 1.30, "Cost": 274.40,
    },
)


def transaction(overrides=None, *, base=0):
    values = dict(TRANSACTIONS[base])
    values.update(overrides or {})
    return values


def transactions_xlsx(*, rows=TRANSACTIONS, headers=HEADERS, header_row=7,
                      column_order=None, formats=None, sheet_name="Transactions",
                      epoch_1904=False, configure=None):
    """Write native Excel date/numeric cells without any real invoice data."""
    workbook = Workbook()
    if epoch_1904:
        workbook.epoch = CALENDAR_MAC_1904
    sheet = workbook.active
    sheet.title = sheet_name
    order = tuple(range(len(HEADERS))) if column_order is None else tuple(column_order)
    field_formats = {"Date": "dd/mm/yyyy hh:mm", "Quantity": "0.00",
                     "Unit Price": "0.0#####", "Cost": "0.00"}
    field_formats.update(formats or {})
    for column, index in enumerate(order, start=1):
        sheet.cell(header_row, column, headers[index])
    for offset, values in enumerate(rows, start=1):
        for column, index in enumerate(order, start=1):
            key = HEADERS[index]
            cell = sheet.cell(header_row + offset, column,
                              None if values is None else values.get(key))
            if key in field_formats:
                cell.number_format = field_formats[key]
    if configure:
        configure(workbook)
    stream = BytesIO()
    workbook.save(stream)
    workbook.close()
    return stream.getvalue()


def changed_zip(content, *, replace=None, append=(), rename=None):
    """Alter synthetic ZIP members for real archive-validation tests."""
    replacement = replace or {}
    renamed = rename or {}
    stream = BytesIO()
    with ZipFile(BytesIO(content)) as original, ZipFile(stream, "w", ZIP_DEFLATED) as output:
        for member in original.infolist():
            output.writestr(renamed.get(member.filename, member.filename),
                            replacement.get(member.filename, original.read(member.filename)))
        for name, data in append:
            output.writestr(name, data)
    return stream.getvalue()
