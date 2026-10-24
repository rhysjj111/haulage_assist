"""Read The Fuel Store's Transactions Excel export into proposed fuel rows.

Only source data is returned. This module neither looks up trucks nor saves
records, calls another service, evaluates formulas, or writes document files.
The supplier profile is an explicit caller choice; Card Type is a network label
and never determines the app's fuel-card supplier.
"""

from dataclasses import dataclass
from datetime import date, datetime, time
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP, localcontext
from io import BytesIO
import re
import warnings
from zipfile import BadZipFile, ZipFile


MAX_XLSX_BYTES = 10 * 1024 * 1024
MAX_EXPANDED_BYTES = 32 * 1024 * 1024
MAX_ZIP_MEMBERS = 1000
MAX_ROWS = 10_000
MAX_COLUMNS = 64
MAX_SHEETS = 10
MAX_HEADER_ROWS = 50
_CENT = Decimal("0.01")
_MAX_NUMBER = Decimal("999999999999.99")
_MAIN_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_WORKBOOK_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"
_WORKBOOK_TYPES = frozenset({
    _WORKBOOK_TYPE,
    "application/vnd.openxmlformats-officedocument.spreadsheetml.template.main+xml",
    "application/vnd.ms-excel.sheet.macroEnabled.main+xml",
    "application/vnd.ms-excel.template.macroEnabled.main+xml",
})
_HEADERS = (
    "Cost Centre", "Card No", "Date", "Card VRN", "Driver", "Tran VRN", "Odo",
    "Card Type", "Product", "Site No", "Site Name", "Quantity", "Unit Price", "Cost",
)
_DIESEL_PRODUCTS = frozenset({
    "DIESEL", "KEYFUELS DIESEL", "UK FUELS DIESEL", "UK FUELS DIESEL (TP)",
    "ESSO DIESEL", "DIESEL (SHELL)",
})
_TWO_DECIMAL_FORMATS = frozenset({
    "0.00", "#,##0.00", "£0.00", "£#,##0.00", '"£"0.00', '"£"#,##0.00',
    "[$£-809]0.00", "[$£-809]#,##0.00", "[$£-en-GB]0.00", "[$£-en-GB]#,##0.00",
})
_MESSAGES = {
    "invalid_workbook": "The spreadsheet could not be read. Use the original Transactions .xlsx export.",
    "too_large": "The spreadsheet is too large or contains too many rows, columns or workbook parts.",
    "unsupported_layout": "Use a Transactions export with the expected columns and one visible worksheet.",
    "unsupported_precision": "Quantity and Cost must use the export's two-decimal number format.",
    "invalid_transaction": "A transaction row is incomplete or has an unsupported date or number. Check the original export.",
    "formula_not_supported": "Formula cells are not supported. Use the original exported transaction values.",
    "encrypted_workbook": "Protected or encrypted spreadsheets are not supported. Use an unprotected .xlsx export.",
    "reader_unavailable": "The spreadsheet reader is unavailable. Install the app requirements and try again.",
}


class FuelXLSXError(ValueError):
    """A fixed safe error, without worksheet values or parser exception text."""

    def __init__(self, code):
        self.code = code
        super().__init__(_MESSAGES[code])


@dataclass(frozen=True)
class FuelTransaction:
    source_row: int
    transaction_date: date
    transaction_time: time | None
    transaction_registration: str | None
    card_registration: str | None
    card_type: str
    product: str
    quantity: Decimal
    cost: Decimal
    unit_price: Decimal
    issues: tuple[str, ...]


@dataclass(frozen=True)
class FuelTransactionsReport:
    sheet_name: str
    header_row: int
    rows: tuple[FuelTransaction, ...]
    source_format: str = "fuel-store-transactions-xlsx-v1"
    fuel_card_name: str = "The Fuel Store"

    @property
    def calculated_total_quantity(self):
        with localcontext() as context:
            context.prec = 50
            return sum((row.quantity for row in self.rows), Decimal("0.00"))

    @property
    def calculated_total_cost(self):
        with localcontext() as context:
            context.prec = 50
            return sum((row.cost for row in self.rows), Decimal("0.00"))


def _normal(value):
    return " ".join(value.split()).casefold() if isinstance(value, str) else None


def _archive_check(content):
    if not isinstance(content, bytes):
        raise FuelXLSXError("invalid_workbook")
    if len(content) > MAX_XLSX_BYTES:
        raise FuelXLSXError("too_large")
    if not content.startswith(b"PK\x03\x04"):
        raise FuelXLSXError("invalid_workbook")
    try:
        with ZipFile(BytesIO(content)) as archive:
            parts = archive.infolist()
            if (len(parts) > MAX_ZIP_MEMBERS
                    or sum(part.file_size for part in parts) > MAX_EXPANDED_BYTES):
                raise FuelXLSXError("too_large")
            names = [part.filename for part in parts]
            if len(set(names)) != len(names):
                raise FuelXLSXError("invalid_workbook")
            if any(part.flag_bits & 1 for part in parts):
                raise FuelXLSXError("encrypted_workbook")
            if any(name.lower().endswith("vbaproject.bin")
                   or name.lower().startswith("xl/externallinks/") for name in names):
                raise FuelXLSXError("unsupported_layout")
    except FuelXLSXError:
        raise
    except (BadZipFile, OSError, ValueError, RuntimeError):
        raise FuelXLSXError("invalid_workbook") from None


def _worksheet_shape_check(content, iterparse, fromstring):
    # Do not trust the worksheet's declared dimensions: a smaller declaration
    # must not hide trailing transactions outside our bounded read window.
    with ZipFile(BytesIO(content)) as archive:
        manifest = fromstring(archive.read("[Content_Types].xml"))
        workbook_parts = [part for part in manifest if part.get("ContentType") in _WORKBOOK_TYPES]
        if (len(workbook_parts) != 1
                or workbook_parts[0].get("ContentType") != _WORKBOOK_TYPE
                or workbook_parts[0].get("PartName") != "/xl/workbook.xml"):
            raise FuelXLSXError("unsupported_layout")
        workbook = fromstring(archive.read("xl/workbook.xml"))
        sheets = workbook.findall(f"{{{_MAIN_NS}}}sheets/{{{_MAIN_NS}}}sheet")
        if len(sheets) > MAX_SHEETS:
            raise FuelXLSXError("too_large")
        if len(sheets) != 1:
            raise FuelXLSXError("unsupported_layout")
        relationships = fromstring(archive.read("xl/_rels/workbook.xml.rels"))
        by_id = {relationship.get("Id"): relationship for relationship in relationships}
        if len(by_id) != len(relationships) or None in by_id:
            raise FuelXLSXError("unsupported_layout")
        names = []
        for sheet in sheets:
            relationship = by_id.get(sheet.get(f"{{{_REL_NS}}}id"))
            if relationship is None or relationship.get("Type") != _REL_NS + "/worksheet":
                raise FuelXLSXError("unsupported_layout")
            target = relationship.get("Target", "")
            # The supported export keeps worksheet parts in this folder. Fail
            # closed on other relationship targets rather than read a sheet
            # whose physical cells were not checked below.
            if (relationship.get("TargetMode") == "External"
                    or not re.fullmatch(r"(?:/xl/)?worksheets/[^/]+\.xml", target)):
                raise FuelXLSXError("unsupported_layout")
            names.append(target.lstrip("/") if target.startswith("/") else "xl/" + target)
        if not names or len(set(names)) != len(names):
            raise FuelXLSXError("unsupported_layout")
        for name in names:
            current_row, current_column = 0, 0
            with archive.open(name) as source:
                for event, element in iterparse(source, events=("start", "end")):
                    if event == "end":
                        element.clear()
                        continue
                    kind = element.tag.rsplit("}", 1)[-1]
                    if kind == "row":
                        value = element.get("r", "")
                        if not re.fullmatch(r"[0-9]{1,8}", value):
                            raise FuelXLSXError("unsupported_layout")
                        row_number = int(value)
                        if not 1 <= row_number <= MAX_ROWS + MAX_HEADER_ROWS:
                            raise FuelXLSXError("too_large")
                        if row_number <= current_row:
                            raise FuelXLSXError("unsupported_layout")
                        current_row, current_column = row_number, 0
                    elif kind == "c":
                        match = re.fullmatch(r"([A-Z]{1,3})([0-9]{1,8})", element.get("r", ""))
                        if not match:
                            raise FuelXLSXError("unsupported_layout")
                        column = 0
                        for letter in match[1]:
                            column = column * 26 + ord(letter) - ord("A") + 1
                        if column > MAX_COLUMNS or not 1 <= int(match[2]) <= MAX_ROWS + MAX_HEADER_ROWS:
                            raise FuelXLSXError("too_large")
                        if int(match[2]) != current_row or column <= current_column:
                            raise FuelXLSXError("unsupported_layout")
                        current_column = column


def _text(cell, *, optional=False):
    value = cell.value
    if optional and (value is None or isinstance(value, str) and not value.strip()):
        return None
    if not isinstance(value, str) or not value.strip() or len(value) > 200:
        raise FuelXLSXError("invalid_transaction")
    return " ".join(value.split())


def _number(cell):
    value = cell.value
    if isinstance(value, bool) or not isinstance(value, (int, float, Decimal)):
        raise FuelXLSXError("invalid_transaction")
    try:
        result = Decimal(str(value))
        if not result.is_finite() or abs(result) > _MAX_NUMBER:
            raise FuelXLSXError("invalid_transaction")
        return result
    except (InvalidOperation, ValueError):
        raise FuelXLSXError("invalid_transaction") from None


def _display_amount(cell, value):
    # This export displays Quantity and Cost to two decimals even when its raw
    # numeric cells contain binary floating-point noise. Follow that explicit
    # display precision; do not guess a rounding policy for other formats.
    if cell.number_format not in _TWO_DECIMAL_FORMATS:
        raise FuelXLSXError("unsupported_precision")
    return value.quantize(_CENT, rounding=ROUND_HALF_UP)


def _transaction(cells, columns, source_row):
    if any(cell.data_type == "f" for cell in cells):
        raise FuelXLSXError("formula_not_supported")
    def cell(label):
        return cells[columns[_normal(label)]]
    occurred = cell("Date").value
    if isinstance(occurred, datetime) and occurred.tzinfo is None:
        on_date, at_time = occurred.date(), occurred.time()
    elif isinstance(occurred, date) and not isinstance(occurred, datetime):
        on_date, at_time = occurred, None
    else:
        raise FuelXLSXError("invalid_transaction")
    registration = _text(cell("Tran VRN"), optional=True)
    card_registration = _text(cell("Card VRN"), optional=True)
    card_type, product = _text(cell("Card Type")), _text(cell("Product"))
    raw_quantity, raw_cost, unit_price = (_number(cell(label))
                                          for label in ("Quantity", "Cost", "Unit Price"))
    quantity = _display_amount(cell("Quantity"), raw_quantity)
    cost = _display_amount(cell("Cost"), raw_cost)
    issues = []
    if not registration:
        issues.append("missing_transaction_registration")
    elif (card_registration
          and "".join(registration.split()).upper() != "".join(card_registration.split()).upper()):
        issues.append("registration_mismatch")
    if product.upper() not in _DIESEL_PRODUCTS:
        issues.append("unsupported_product")
    if quantity <= 0:
        issues.append("non_positive_quantity")
    if cost <= 0:
        issues.append("non_positive_cost")
    if (raw_quantity * unit_price).quantize(_CENT, rounding=ROUND_HALF_UP) != cost:
        issues.append("cost_price_mismatch")
    return FuelTransaction(
        source_row=source_row, transaction_date=on_date, transaction_time=at_time,
        transaction_registration=registration.upper() if registration else None,
        card_registration=card_registration.upper() if card_registration else None,
        card_type=card_type, product=product, quantity=quantity, cost=cost,
        unit_price=unit_price, issues=tuple(issues),
    )


def _sheet_report(workbook):
    if len(workbook.sheetnames) > MAX_SHEETS:
        raise FuelXLSXError("too_large")
    if len(workbook.sheetnames) != 1:
        raise FuelXLSXError("unsupported_layout")
    sheet = workbook[workbook.sheetnames[0]]
    if _normal(sheet.title) != "transactions" or sheet.sheet_state != "visible":
        raise FuelXLSXError("unsupported_layout")
    if (sheet.max_column is not None and sheet.max_column > MAX_COLUMNS
            or sheet.max_row is not None and sheet.max_row > MAX_ROWS + MAX_HEADER_ROWS):
        raise FuelXLSXError("too_large")
    expected = {_normal(header) for header in _HEADERS}
    header_row, columns = None, None
    for index, cells in enumerate(sheet.iter_rows(min_row=1, max_row=MAX_HEADER_ROWS,
                                                 max_col=MAX_COLUMNS), start=1):
        nonempty = [(i, _normal(cell.value)) for i, cell in enumerate(cells) if cell.value is not None]
        if len(nonempty) == len(expected) and {label for _, label in nonempty} == expected:
            if header_row is not None:
                raise FuelXLSXError("unsupported_layout")
            header_row, columns = index, {label: i for i, label in nonempty}
    if header_row is None:
        raise FuelXLSXError("unsupported_layout")
    rows = []
    for index, cells in enumerate(sheet.iter_rows(min_row=header_row + 1,
                                                 max_row=MAX_ROWS + MAX_HEADER_ROWS + 1,
                                                 max_col=MAX_COLUMNS), start=header_row + 1):
        if not any(cell.value is not None for cell in cells):
            continue
        if any(cell.value is not None for i, cell in enumerate(cells) if i not in columns.values()):
            raise FuelXLSXError("unsupported_layout")
        if len(rows) >= MAX_ROWS:
            raise FuelXLSXError("too_large")
        rows.append(_transaction(cells, columns, index))
    if not rows:
        raise FuelXLSXError("unsupported_layout")
    return FuelTransactionsReport(sheet_name=sheet.title, header_row=header_row, rows=tuple(rows))


def read_fuel_store_transactions_xlsx(content: bytes) -> FuelTransactionsReport:
    """Read the complete supported export or raise a safe :class:`FuelXLSXError`.

    The returned Cost is the source column, without VAT calculations. Flags
    retain unusual rows for later review; they never remove or correct a row.
    ZIP/shape limits bound ordinary inputs, but this direct function is not an
    execution sandbox. A future HTTP upload must also isolate time and memory.
    """
    _archive_check(content)
    try:
        # Require the XML protection used automatically by openpyxl. Imports are
        # lazy so a missing optional reader dependency cannot prevent app boot.
        from defusedxml.ElementTree import fromstring, iterparse
        import openpyxl
        from openpyxl.xml import DEFUSEDXML
    except ImportError:
        raise FuelXLSXError("reader_unavailable") from None
    if not DEFUSEDXML:
        raise FuelXLSXError("reader_unavailable")
    workbook = None
    try:
        with warnings.catch_warnings(), localcontext() as context:
            # openpyxl warnings can include raw cell values, even from ignored
            # identity columns. Keep all parser diagnostics out of app logs.
            warnings.simplefilter("ignore")
            context.prec = 50
            _worksheet_shape_check(content, iterparse, fromstring)
            workbook = openpyxl.load_workbook(BytesIO(content), read_only=True,
                                             data_only=False, keep_links=False)
            return _sheet_report(workbook)
    except FuelXLSXError:
        raise
    except Exception:
        # Library failures can contain arbitrary cell values or document XML.
        # Never include their messages or return a partially parsed transaction.
        raise FuelXLSXError("invalid_workbook") from None
    finally:
        if workbook is not None:
            workbook.close()
